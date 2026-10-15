from __future__ import annotations

import json
import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/roselindsi/Documents/GitHub/video-split")
OUT = ROOT / "docs" / "视频动作图与Episode Identity研究进展报告_20260923.docx"
BUILD = ROOT / ".artifacts" / "episode_identity_progress"
CHART = BUILD / "qwen_blind30_summary.png"
GOLD = ROOT / "results/auditor/interaction_graph_blind30_v2/human/reviewer_1_final.json"
SCORE = ROOT / "results/auditor/interaction_graph_blind30_v2/qwen_5fold_v1/qwen/score/agreement.json"

BLACK = "000000"
INK = "243238"
MUTED = "5E6A70"
HEADER_FILL = "365B66"
PALE_FILL = "EEF4F4"
ALT_FILL = "F7F9F9"
BORDER = "D9D9D9"
ACCENT = "2E7D73"
WARM = "C56B3F"


def set_run_font(run, name="Arial Unicode MS", size=None, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, name, size, bold=False, color=BLACK):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_borders(style_or_paragraph):
    if hasattr(style_or_paragraph, "_element") and style_or_paragraph._element.tag == qn("w:style"):
        p_pr = style_or_paragraph._element.get_or_add_pPr()
    else:
        p_pr = style_or_paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8.5, color=MUTED)


def add_rich_paragraph(doc, parts, style=None, align=None, space_after=6, first_indent=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(21)
    for part in parts:
        if isinstance(part, str):
            text, bold, color, italic = part, False, BLACK, False
        else:
            text = part.get("text", "")
            bold = part.get("bold", False)
            color = part.get("color", BLACK)
            italic = part.get("italic", False)
        run = p.add_run(text)
        set_run_font(run, size=10.8, bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc, text, level=0, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.24 + 0.18 * level)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10.6, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.6)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.6)
    return p


def add_numbered_step(doc, number, title, input_text, process_text, output_text, acceptance_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{number}  {title}")
    set_run_font(r, size=11.5, bold=True, color=INK)
    for label, value in (
        ("输入", input_text),
        ("处理", process_text),
        ("输出", output_text),
        ("验收标准", acceptance_text),
    ):
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Inches(0.20)
        q.paragraph_format.space_after = Pt(2)
        q.paragraph_format.line_spacing = 1.23
        a = q.add_run(f"{label}  ")
        set_run_font(a, size=10.1, bold=True, color=ACCENT)
        b = q.add_run(value)
        set_run_font(b, size=10.1)


def add_table(doc, headers, rows, widths, font_size=8.9, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, HEADER_FILL)
        set_cell_margins(cell, top=115, bottom=115)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(text))
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.width = Inches(widths[i])
            set_cell_margins(cell, top=105, bottom=105)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                set_cell_shading(cell, ALT_FILL)
            p = cell.paragraphs[0]
            if alignments:
                p.alignment = alignments[i]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(text))
            set_run_font(r, size=font_size, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    return table


def add_source_path(doc, label, path):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=9.2, bold=True, color=INK)
    r2 = p.add_run(path)
    set_run_font(r2, name="Menlo", size=8.4, color=MUTED)


def load_stats():
    gold = json.loads(GOLD.read_text(encoding="utf-8"))
    videos = gold["videos"]
    human = {
        "videos": len(videos),
        "events": sum(len(v["events"]) for v in videos),
        "segments": sum(len(v["segments"]) for v in videos),
        "boundaries": sum(max(0, len(v["segments"]) - 1) for v in videos),
        "repeat_segments": sum(
            1 for v in videos for s in v["segments"] if s.get("repeats_within_segment")
        ),
        "multi_segment_videos": sum(1 for v in videos if len(v["segments"]) > 1),
    }
    score = json.loads(SCORE.read_text(encoding="utf-8"))
    return human, score


def make_chart(human, score):
    BUILD.mkdir(parents=True, exist_ok=True)
    width, height = 2200, 880
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/PingFang.ttc"
    bold_path = "/System/Library/Fonts/STHeiti Medium.ttc"
    title_font = ImageFont.truetype(bold_path, 46)
    panel_font = ImageFont.truetype(bold_path, 34)
    label_font = ImageFont.truetype(font_path, 27)
    small_font = ImageFont.truetype(font_path, 24)
    number_font = ImageFont.truetype(bold_path, 29)

    ink = "#243238"
    muted = "#68757A"
    grid = "#E1E6E6"
    green = "#2E7D73"
    orange = "#C56B3F"

    def center_text(x, y, text_value, font, fill=ink):
        box = draw.textbbox((0, 0), text_value, font=font)
        draw.text((x - (box[2] - box[0]) / 2, y), text_value, font=font, fill=fill)

    center_text(width / 2, 20, "30 条盲标中 Qwen 的核心误差", title_font)

    # Left panel: segment and boundary counts.
    left = (90, 120, 1030, 785)
    center_text((left[0] + left[2]) / 2, 105, "人工与 Qwen 的切分数量", panel_font)
    plot_left, plot_top, plot_right, plot_bottom = 170, 200, 960, 690
    max_val = 140
    for tick in (0, 35, 70, 105, 140):
        y = plot_bottom - (tick / max_val) * (plot_bottom - plot_top)
        draw.line((plot_left, y, plot_right, y), fill=grid, width=2)
        draw.text((105, y - 14), str(tick), font=small_font, fill=muted)
    draw.line((plot_left, plot_top, plot_left, plot_bottom), fill=ink, width=2)
    draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill=ink, width=2)

    labels = ["时间段", "边界"]
    human_vals = [121, 93]
    qwen_vals = [80, 52]
    centers = [390, 740]
    bar_w = 92
    gap = 18
    for cx, label, hv, qv in zip(centers, labels, human_vals, qwen_vals):
        for value, color, x in ((hv, green, cx - bar_w - gap / 2), (qv, orange, cx + gap / 2)):
            top = plot_bottom - (value / max_val) * (plot_bottom - plot_top)
            draw.rounded_rectangle((x, top, x + bar_w, plot_bottom), radius=6, fill=color)
            center_text(x + bar_w / 2, top - 38, str(value), number_font)
        center_text(cx, plot_bottom + 22, label, label_font)
    draw.rectangle((705, 170, 730, 195), fill=green)
    draw.text((742, 163), "人工", font=small_font, fill=ink)
    draw.rectangle((820, 170, 845, 195), fill=orange)
    draw.text((857, 163), "Qwen", font=small_font, fill=ink)

    # Right panel: same-action/new-instance funnel.
    rleft, rright = 1170, 2125
    center_text((rleft + rright) / 2, 105, "同类新实例识别漏斗", panel_font)
    axis_x, axis_right = 1540, 2050
    axis_top, axis_bottom = 205, 690
    max_funnel = 55
    for tick in (0, 10, 20, 30, 40, 50):
        x = axis_x + (tick / max_funnel) * (axis_right - axis_x)
        draw.line((x, axis_top, x, axis_bottom), fill=grid, width=2)
        center_text(x, axis_bottom + 20, str(tick), small_font, fill=muted)
    draw.line((axis_x, axis_top, axis_x, axis_bottom), fill=ink, width=2)
    funnel_labels = ["人工同类新实例", "Qwen 预测", "1 秒内定位", "定位且关系正确"]
    vals = [51, 6, 3, 1]
    colors = [green, "#799E98", "#D6A07B", orange]
    ys = [245, 355, 465, 575]
    for label, value, color, y in zip(funnel_labels, vals, colors, ys):
        box = draw.textbbox((0, 0), label, font=label_font)
        draw.text((axis_x - 26 - (box[2] - box[0]), y + 10), label, font=label_font, fill=ink)
        bar_right = axis_x + max(12, (value / max_funnel) * (axis_right - axis_x))
        draw.rounded_rectangle((axis_x, y, bar_right, y + 58), radius=6, fill=color)
        draw.text((bar_right + 14, y + 10), str(value), font=number_font, fill=ink)

    center_text(width / 2, 820, "模型能看到重复，但很少把同类动作的重新开始切成新 episode", label_font, fill=muted)
    image.save(CHART, quality=95)


def build_document():
    human, score = load_stats()
    make_chart(human, score)
    OUT.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    props = doc.core_properties
    props.title = "视频动作图与 Episode Identity 研究进展报告"
    props.subject = "Taxonomy free interaction graph research progress"
    props.author = "项目研究记录"
    props.keywords = "episode identity, interaction graph, video understanding, timeline annotation"

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.80)
    section.right_margin = Inches(0.80)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.30)

    styles = doc.styles
    set_style_font(styles["Normal"], "Arial Unicode MS", 10.8, False, BLACK)
    styles["Normal"].paragraph_format.line_spacing = 1.35
    styles["Normal"].paragraph_format.space_after = Pt(6)
    set_style_font(styles["Title"], "Arial Unicode MS", 24, True, BLACK)
    styles["Title"].paragraph_format.space_after = Pt(14)
    remove_paragraph_borders(styles["Title"])
    set_style_font(styles["Subtitle"], "Arial Unicode MS", 12, False, MUTED)
    set_style_font(styles["Heading 1"], "Arial Unicode MS", 16, True, BLACK)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    set_style_font(styles["Heading 2"], "Arial Unicode MS", 12.5, True, BLACK)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Bullet 2"):
        set_style_font(styles[style_name], "Arial Unicode MS", 10.6, False, BLACK)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.paragraph_format.space_after = Pt(0)
    hr = header_p.add_run("视频动作图与 Episode Identity 研究进展")
    set_run_font(hr, size=8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    # Cover and opening conclusion
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_borders(title)
    tr = title.add_run("视频动作图与 Episode Identity 研究进展报告")
    set_run_font(tr, size=24, bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Taxonomy Free Interaction Graphs")
    set_run_font(sr, size=12, color=MUTED)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(18)
    mr = meta.add_run("内部研究进展报告  |  2026 年 9 月 23 日  |  资料截至 2026 年 9 月 22 日")
    set_run_font(mr, size=10, color=MUTED)
    for _ in range(3):
        doc.add_paragraph()
    lead = doc.add_paragraph()
    lead.paragraph_format.left_indent = Inches(0.55)
    lead.paragraph_format.right_indent = Inches(0.55)
    lead.paragraph_format.line_spacing = 1.45
    lead.paragraph_format.space_after = Pt(8)
    r = lead.add_run(
        "本研究要回答一个具体问题：视频模型能否知道眼前动作仍属于同一次 episode，"
        "还是同类动作已经重新开始，或操作者已经转向另一个任务。"
    )
    set_run_font(r, size=12, bold=True, color=INK)
    lead2 = doc.add_paragraph()
    lead2.paragraph_format.left_indent = Inches(0.55)
    lead2.paragraph_format.right_indent = Inches(0.55)
    lead2.paragraph_format.line_spacing = 1.45
    r = lead2.add_run(
        "目前已经完成图目标定义、完整时间轴标注工具、60 条开发数据的图发布、"
        "一批新的 30 条盲标，以及 Qwen 的五折盲测。现有结果给出了清楚的负面诊断："
        "Qwen 能看到重复动作，但经常把同类新实例吞进同一个持续 episode。"
        "下一阶段的关键不是立即扩大自动标注，而是先完成新版时间轴合同的双人人工验证，"
        "再把模型任务拆成边界定位、边界关系分类和节点复用三个子问题。"
    )
    set_run_font(r, size=11, color=INK)

    doc.add_page_break()

    doc.add_heading("执行摘要", level=1)
    add_bullet(doc, "研究对象已经从固定动作 taxonomy 转为录像内部的匿名事件节点、episode 实例和时间边界。", bold_lead="研究对象")
    add_bullet(doc, "图构建管线已经可运行：60 条开发视频形成 134 个动作节点、45 个 procedure groups 和 341 次细粒度边出现。", bold_lead="图构建管线")
    add_bullet(doc, "旧候选点问卷的双人主关系一致率仅为 36.7%，直接暴露了目标颗粒度和时间锚定不统一的问题。", bold_lead="旧问卷结果")
    add_bullet(doc, "改用完整时间轴后，人工完成新的 30 条盲标，共 127 段、97 个边界；但当前仍只有一位人工标注者。", bold_lead="新版盲标")
    add_bullet(doc, "Qwen 五折盲测的 1 秒边界 F1 为 0.138；51 个同类新实例边界中，仅 3 个被定位，且只有 1 个关系判断正确。", bold_lead="模型结果")
    add_bullet(doc, "因此，现在可以报告模型的明确失败模式，但不能把当前图合同称为已通过人际一致性验证，也不能把本轮模型结果当成最终确认性结论。", bold_lead="结论边界")

    doc.add_heading("目前最重要的判断", level=2)
    add_rich_paragraph(doc, [
        {"text": "这条研究线已经越过“能否建图”的阶段，进入“图的 episode 语义是否可复现，以及模型为什么识别失败”的阶段。", "bold": True},
        " 数据与工具已经足以开展严谨评估；真正欠缺的是新版合同的第二位人工标注者，以及冻结后的全新确认集。",
    ])

    doc.add_heading("一 研究课题", level=1)
    add_rich_paragraph(doc, [
        "现实视频里的动作往往不是一次干净的类别切换。擦桌子会来回多次，拆风扇会连续拆头部和支撑杆，拿出下水道塞子后还会冲洗并放回。",
        {"text": "研究问题不是单纯识别动作名称，而是识别同一个行为目标在时间上的身份。", "bold": True},
    ], first_indent=True)
    add_rich_paragraph(doc, [
        "我们把这个问题称为 ",
        {"text": "episode identity", "bold": True, "italic": True},
        "：模型需要区分同一实例内部的重复、同类型的新实例，以及真正的任务转移。图只是把这种判断显式化，使错误能够被定位和计数。",
    ], first_indent=True)

    add_table(
        doc,
        ["时间上的情况", "时间轴标注", "图上的含义", "直观例子"],
        [
            ["同一次 episode 内继续或反复", "不切边界；可标段内重复", "保持同一实例；不产生边", "同一块桌面来回擦"],
            ["同类动作重新开始", "切边界并复用同一事件节点", "A → A；新 instance", "一个螺丝完成后再上下一颗"],
            ["目标转为另一种动作", "切边界并建立或复用另一节点", "A → B", "缠线后开始解线；炒菜后去扫地"],
            ["同一流程内的不同步骤", "保留不同局部节点", "节点间为组内转移", "拆风扇头与逐节拆杆同属拆卸风扇"],
        ],
        [1.55, 1.65, 1.35, 2.15],
        font_size=8.7,
    )
    add_rich_paragraph(doc, [
        {"text": "关键修正：", "bold": True},
        "普通的持续重复不再被当作 self loop。只有上一实例已经结束或重置、同类目标重新开始时，才记录 A → A。这样才能让“内部重复”和“同类新实例”成为可检验的不同对象。",
    ])

    doc.add_page_break()

    doc.add_heading("二 图表示与颗粒度", level=1)
    doc.add_heading("局部事件节点", level=2)
    add_rich_paragraph(doc, [
        "节点表示一条录像内部可复用的、可观察的即时结果，而不是预先规定的全局动作类别。事件字母 A、B、C 只在本录像内有效。"
        "同一物品可以对应不同节点，例如“缠线”和“解线”；不同物品实例也可以复用同一节点，例如连续处理多根等价电线。",
    ], first_indent=True)

    doc.add_heading("Episode 实例", level=2)
    add_rich_paragraph(doc, [
        "episode 是一次目标追求。方向反复、重新抓握、短暂停顿和中途松手都可能发生在同一个实例内。"
        "只有结果可见地完成、尝试终止或重置后又重新开始，才建立新的 instance。",
    ], first_indent=True)

    doc.add_heading("Procedure group", level=2)
    add_rich_paragraph(doc, [
        "procedure_group 是局部事件节点之上的父层，用来表达多个步骤服务于同一直接流程。分组同时要求目标、对象家族和结果方向兼容。"
        "它不会删除细粒度节点，也不会把组内不同步骤误写成完全相同的动作。",
    ], first_indent=True)
    add_bullet(doc, "`旋松并取下风扇头部` 与 `逐节拆下风扇支撑杆` 可以同属 `拆卸风扇`。")
    add_bullet(doc, "`前往另一房间取凳子并返回` 属于 `取用辅助物品`，不是拆风扇本身。")
    add_bullet(doc, "`取下手机壳` 与 `装回手机壳` 保留为不同局部状态，但可归入更宽的 `装卸手机壳` 流程组。")

    doc.add_heading("特殊内容与不可见内容", level=2)
    add_rich_paragraph(doc, [
        "`特殊事件`、`看不到内容`、`idle` 等状态保留在源时间轴中以便追溯，但不进入动作节点和 procedure groups。"
        "若特殊状态位于两个动作之间，不跨越它自动补边，避免凭空制造转移。",
    ], first_indent=True)

    doc.add_heading("三 方法演进与已完成工作", level=1)
    add_bullet(doc, "阶段一  候选点关系问卷。两位标注者在 30 条难例上独立回答 CONTINUE、SAME_ACTION_NEW_INSTANCE、NEW_ACTION 等关系。")
    add_bullet(doc, "阶段二  分歧诊断。主关系仅 11/30 一致，Cohen kappa 为 0.191；时间锚定和目标颗粒度均不统一。")
    add_bullet(doc, "阶段三  完整时间轴标注网站。加入可编辑切点、视频时间轴映射、左右箭头选择窄片段、节点复用、段内重复、草稿持久化和离线包。")
    add_bullet(doc, "阶段四  两批 30 条开发标注。第一批提供大量往返和回环；手部马赛克补集增加装配、线束加工、上架、清洁、运输与协作等内容。")
    add_bullet(doc, "阶段五  规范化与图发布。保留局部节点，增加 procedure_group 层，发布细粒度边、组级边、检查队列、冻结标注和校验和。")
    add_bullet(doc, "阶段六  新的 30 条盲选与人工时间轴。选样不使用动作标签、模型分数或已有预测，且与上一批手部马赛克样本零重叠。")
    add_bullet(doc, "阶段七  Qwen 五折盲测。每折 6 条测试视频，目标视频的人工答案不进入提示；结果用于诊断模型，不替代第二位人工标注者。")

    doc.add_page_break()

    doc.add_heading("四 已有数据与图发布", level=1)
    add_rich_paragraph(doc, [
        "目前的可复查开发发布由两批各 30 条视频组成。两批的行为结构不同：第一批更擅长暴露返回旧节点的循环，马赛克补集更擅长覆盖多步骤工业与日常流程。",
    ])
    add_table(
        doc,
        ["指标", "重复动作批", "马赛克补集", "合计或说明"],
        [
            ["视频数", "30", "30", "60"],
            ["时长", "39.55 分钟", "53.38 分钟", "92.93 分钟"],
            ["动作节点", "65", "69", "134"],
            ["原始时间段", "238", "174", "412"],
            ["多动作视频", "18", "20", "38"],
            ["有 self loop 证据的视频", "12", "22", "口径为重复存在性"],
            ["回到先前动作的视频", "18", "7", "两批拓扑互补"],
            ["动作转移速度", "3.26 次/分钟", "0.96 次/分钟", "不应互相替代"],
        ],
        [2.0, 1.25, 1.25, 2.2],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("Interaction Graph Release v1", level=2)
    add_bullet(doc, "60 条视频、134 个动作节点、45 个 procedure groups；其中 8 个组跨视频复用。")
    add_bullet(doc, "341 次细粒度边出现：161 次 exact state repeat、171 次组内转移、9 次跨组转移。")
    add_bullet(doc, "形成 141 条聚合细粒度边和 62 条聚合组级边；4 条视频出现跨组转移，2 条视频离开后返回原 procedure group。")
    add_bullet(doc, "39 次特殊状态端点被省略，且没有跨特殊状态间隔合成替代边。")
    add_bullet(doc, "9 个节点保留为中等置信度，进入人工复核队列；源时间轴未被改写。")
    add_rich_paragraph(doc, [
        {"text": "证据边界：", "bold": True},
        "该发布是完整、可复现的单一标注者开发版本。它证明建图流程能够工作，但不能被表述为已获得人际一致性。",
    ])

    doc.add_heading("五 新一轮 30 条盲标", level=1)
    add_rich_paragraph(doc, [
        "新批次在观看动作内容之前冻结。它从未进入上一批马赛克补集的 databag 中选取，按马赛克关键帧数量分成低、中、高三个层，每层 10 条，且 30 条来自 30 个不同来源组。",
    ], first_indent=True)
    add_table(
        doc,
        ["项目", "结果"],
        [
            ["样本", "30 条视频；26 条为 120 秒，4 条受源视频长度限制"],
            ["人工事件节点", f"{human['events']} 个录像内局部节点"],
            ["人工时间段", f"{human['segments']} 段"],
            ["人工边界", f"{human['boundaries']} 个"],
            ["包含多个时间段的视频", f"{human['multi_segment_videos']} / 30"],
            ["段内重复标记", f"{human['repeat_segments']} 段"],
            ["标注状态", "第一位人工标注者已完成；第二位人工标注者尚未进行"],
        ],
        [2.0, 4.7],
        font_size=9.2,
    )
    add_rich_paragraph(doc, [
        "这 30 条比开发发布更接近真实研究测试，但当前人工答案仍是单人参考。Qwen 与这份参考的比较能够发现模型错误，不能估计人类可复现性。",
    ])

    doc.add_page_break()

    doc.add_heading("六 Qwen 五折盲测", level=1)
    doc.add_heading("实验设计", level=2)
    add_bullet(doc, "模型为 Qwen3 VL 8B Instruct；视频按 1 fps 采样；温度为 0。")
    add_bullet(doc, "五折交叉测试，每折保留 6 条测试视频，每个马赛克层级 2 条。")
    add_bullet(doc, "测试视频的人工事件、边界、描述和备注均不进入该视频的提示。")
    add_bullet(doc, "模型先独立建立时间轴，再进行第二遍检查；输出与人工网站相同的 events 和 segments JSON。")
    add_bullet(doc, "30 条中 28 条产生有效结构化输出；2 条因不断生成同义事件节点并在 JSON 完成前截断而失败。")

    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_before = Pt(5)
    pic_p.paragraph_format.space_after = Pt(3)
    pic_p.add_run().add_picture(str(CHART), width=Inches(6.55))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cr = cap.add_run("图 1  Qwen 的主要问题不是过切，而是漏掉同类新实例边界")
    set_run_font(cr, size=9, color=MUTED)

    add_table(
        doc,
        ["指标", "结果", "如何解读"],
        [
            ["有效输出", "28 / 30", "两条结构化输出失败"],
            ["边界 F1 0.25 秒", "0.055", "严格时间定位几乎失败"],
            ["边界 F1 0.5 秒", "0.069", "扩大容差仍很低"],
            ["边界 F1 1.0 秒", "0.138", "主要指标明显不达标"],
            ["边界 F1 2.0 秒", "0.234", "并非只差一两帧"],
            ["匹配边界关系一致率", "0.800  8 / 10", "仅基于 10 个已匹配边界"],
            ["时间分区 ARI 中位数", "0.580", "会被单节点视频虚高"],
            ["完整归一化路径", "5 / 28", "整条事件路径很少完全一致"],
            ["段内重复一致率", "0.593", "重复范围过宽"],
            ["段内重复 kappa", "0.253", "超过偶然的一致性有限"],
        ],
        [2.0, 1.35, 3.35],
        font_size=8.6,
    )

    doc.add_page_break()

    doc.add_heading("七 结果解释", level=1)
    doc.add_heading("模型暴露出的具体失败模式", level=2)
    add_rich_paragraph(doc, [
        "在 28 条有效输出中，人工有 121 段和 93 个边界，Qwen 只有 80 段和 52 个边界。16 条视频明显欠切，11 条人工多段视频被压成单段。",
    ], first_indent=True)
    add_rich_paragraph(doc, [
        "人工边界中有 51 个 `SAME_ACTION_NEW_INSTANCE` 和 42 个 `NEW_ACTION`。Qwen 只预测 6 个同类新实例边界。"
        "在 1 秒容差内，它只找到 3 / 51 个同类新实例边界，其中只有 1 个同时把关系判断正确。",
    ], first_indent=True)
    add_rich_paragraph(doc, [
        {"text": "核心结论是：", "bold": True},
        "Qwen 能识别大量表面上的重复动作，但经常把“上一轮已经结束、同类动作重新开始”继续解释为同一段内部重复。"
        "这正是研究问题要检测的 episode identity 失败。",
    ])

    doc.add_heading("为什么不能只看关系一致率或 ARI", level=2)
    add_bullet(doc, "0.800 的关系一致率来自 10 个已经对齐的边界。模型先漏掉了大部分边界，因此这个分母太小，不能代表整体表现。")
    add_bullet(doc, "单节点视频即使完全漏掉多个同类实例边界，事件类型分区仍可能得到很高的 ARI。")
    add_bullet(doc, "边界召回与同类新实例召回是当前研究问题更直接的指标；分区指标只能作为补充。")

    doc.add_heading("现在可以说什么", level=2)
    add_table(
        doc,
        ["可以支持的结论", "暂时不能支持的结论"],
        [
            ["完整时间轴图可以把 episode identity 错误显式化。", "当前图合同已通过双人人工验证。"],
            ["Qwen 在这批视频上存在系统性欠切。", "Qwen 可以替代第二位人工标注者。"],
            ["同类新实例是当前最弱的关系。", "0.800 的匹配关系一致率表示整体关系识别良好。"],
            ["需要把整体生成任务拆成多个可诊断步骤。", "本轮结果已经是冻结后的最终确认性测试。"],
        ],
        [3.35, 3.35],
        font_size=9.0,
    )
    add_rich_paragraph(doc, [
        "本轮在观察到格式失败后修复了原始文本保存、未使用事件清理和第二遍重试逻辑。修复没有人工补写边界，但因为工程规则是在看到失败后调整的，",
        {"text": "本轮应报告为开发性评估，而不是最终确认性测试。", "bold": True},
    ])

    doc.add_heading("八 当前证据成熟度", level=1)
    add_table(
        doc,
        ["研究组成", "当前状态", "证据强度", "主要缺口"],
        [
            ["图目标与标注规则", "已形成版本化合同", "中", "新版完整时间轴尚未通过双人复现"],
            ["标注工具", "在线与离线版可用", "较强", "需继续锁定导出格式和版本"],
            ["60 条开发图发布", "已冻结并带校验和", "较强", "单一标注者；9 个中置信节点待复核"],
            ["新 30 条人工盲标", "第一位已完成", "中", "缺第二位独立人工标注"],
            ["Qwen 五折评估", "已完成开发性测试", "中", "2 条输出失败；尚未在全新确认集复现"],
            ["Episode identity 结论", "发现明确候选失败模式", "中", "需人际一致性和确认性模型测试共同加固"],
            ["并行动作与遮挡", "合同已有原则", "偏弱", "缺足量正例与专门验证"],
        ],
        [1.55, 1.60, 0.85, 2.7],
        font_size=8.5,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_page_break()

    doc.add_heading("九 下一步工作计划", level=1)
    add_rich_paragraph(doc, [
        "以下步骤按依赖顺序排列。前 3 步决定标注目标是否可复现，后 4 步才进入模型改进与确认性结论。",
    ])
    add_numbered_step(
        doc, 1, "冻结新版合同与工具",
        "现有 target contract、procedure group contract、完整时间轴网站、评分脚本和导出格式。",
        "统一版本号，记录校验和；把内部重复不出边、同类新实例 A→A、新动作 A→B 的规则写入同一手册。",
        "一个不依赖口头补充的冻结标注包。",
        "第二位标注者仅凭包内材料即可开始，且工具输出可被评分器直接读取。",
    )
    add_numbered_step(
        doc, 2, "完成第二位人工独立标注",
        "已经冻结的 30 条 whole timeline 双人验证包。",
        "第二位标注者独立观看完整视频，不接触第一位答案、Qwen 输出或案例讨论。",
        "第二份完整 events 和 segments 导出。",
        "30 条全部存在且无多余条目；保留原始结果，不在评分前改答案。",
    )
    add_numbered_step(
        doc, 3, "评分并裁决合同",
        "两位人工时间轴和冻结评分阈值。",
        "计算 0.25、0.5、1 秒边界 F1、1 秒内关系一致率、时间分区 ARI 和段内重复一致率；评分后再逐条裁决。",
        "人际一致性报告、分歧清单、裁决记录和合同 v2。",
        "1 秒边界 F1 ≥ 0.80、匹配关系一致率 ≥ 0.80、ARI 中位数 ≥ 0.80、段内重复一致率 ≥ 0.80；若未通过，先修规则再重标受影响样本。",
    )
    add_numbered_step(
        doc, 4, "把模型任务拆成三步",
        "人工时间轴、现有 Qwen 失败案例和视频。",
        "分别训练或提示模型完成候选边界定位、CONTINUE/SAME_ACTION_NEW_INSTANCE/NEW_ACTION 分类、节点复用与 episode 链接。",
        "三个可单独评分的模型输出，而非一次生成整张图。",
        "每一步都能独立统计召回、精确率和错误类型；边界漏检不再被关系指标掩盖。",
    )

    doc.add_page_break()

    doc.add_heading("九 下一步工作计划 续", level=1)
    add_numbered_step(
        doc, 5, "完成开发性误差修复",
        "三步模型在现有 30 条上的输出。",
        "重点处理同类新实例欠召回、重复范围过宽、同义节点爆炸和 JSON 截断；建立按视频复杂度分层的误差表。",
        "冻结的模型代码、提示、解码参数和格式验证器。",
        "所有 30 条输出可解析；任何后处理只做预先定义的格式规范化，不人工合成边界。",
    )
    add_numbered_step(
        doc, 6, "预注册并运行全新确认集",
        "完全未见的新 30 条视频和已冻结模型。",
        "在打开人工答案前写定主要指标、容差、失败处理和成功阈值；随后一次性运行并锁定结果。",
        "确认性测试报告。",
        "无样本泄漏、无看结果后修规则；同时报告边界召回、同类新实例召回和失败输出覆盖率。",
    )
    add_numbered_step(
        doc, 7, "扩大图并回答主研究问题",
        "通过双人合同的人工数据、确认性模型结果和 procedure groups。",
        "比较不同模型、视频复杂度、重复结构和流程组上的 episode identity；分析错误是否集中在外观相似但实例更新的情形。",
        "可用于论文的主表、误差分类和图结构案例。",
        "结论明确区分人工可靠性、模型性能和数据分布；所有主张都能回溯到冻结样本与版本。",
    )

    doc.add_heading("推荐的近期优先级", level=2)
    add_rich_paragraph(doc, [
        {"text": "第一优先级是步骤 2 和 3。", "bold": True},
        "在没有新版时间轴双人一致性之前，继续增加模型标签只会把标注定义的不确定性一起放大。"
        "双人验证完成后，再冻结三步模型，下一轮数据才真正具有确认性价值。",
    ])

    doc.add_heading("十 平行的手部马赛克技术线", level=1)
    add_rich_paragraph(doc, [
        "手部马赛克与 episode identity 是相关但不同的研究层。前者解决手部观测是否真实、轨迹是否串号和弱检测是否需要语义过滤；后者解决动作实例在时间上的连续与更新。两条线不应混用指标。",
    ], first_indent=True)
    add_bullet(doc, "人工已确认 52 个持续翻边事件均属于 owner，`52 / 52 owner` 本身没有标注问题。")
    add_bullet(doc, "25.3% 的分母是 87 条长度足以判断的链，应称为 sustained side reversal incidence；它既不是 169 条全体链的比例，也不是严格上界。")
    add_bullet(doc, "原 blur check 缺 tracker id，因此当时无法区分后处理 stitching 串错和 tracker 原生串号；需要 tid 对齐结果才能分解。")
    add_bullet(doc, "Qwen 手与非手语义探针在 88 个 gold 上达到 AUC 0.950；阈值 0.60 时平衡准确率 92.0%。它适合做离线弱观测闸门，但不是 episode identity 的证据。")

    doc.add_page_break()

    doc.add_heading("十一 当前结论", level=1)
    add_rich_paragraph(doc, [
        "当前项目已经拥有清晰的问题定义、能用的完整时间轴标注工具、两批互补开发数据、可复现的图发布和一轮真正按折隐藏目标答案的模型测试。"
        "这些工作足以证明研究问题是可操作的，也足以指出 Qwen 的主要问题不是“看不见重复”，而是“不知道一次重复何时已经结束并重新开始”。",
    ], first_indent=True)
    add_rich_paragraph(doc, [
        "但当前阶段还不能宣称图合同已成为可靠金标准。旧问卷的双人结果失败，新版完整时间轴尚缺第二位人工标注者；模型测试又是在格式修复后完成的开发性评估。"
        "因此最稳妥的研究路线是：先通过新版双人验证冻结目标，再使用分解后的模型完成全新确认集，最后扩大图规模并比较模型对 episode identity 的识别能力。",
    ], first_indent=True)
    add_rich_paragraph(doc, [
        {"text": "一句话总结：", "bold": True},
        "图已经建得出来，失败模式也已经看见；接下来要先证明人能稳定地按同一规则建图，再证明模型能否跨过“内部重复”和“同类新实例”之间的那条边界。",
    ])

    doc.add_heading("附录 关键项目材料", level=1)
    add_source_path(doc, "图目标合同", "docs/interaction_graph_target_contract_v1.md")
    add_source_path(doc, "Procedure group 合同", "docs/interaction_graph_procedure_group_contract_v1.md")
    add_source_path(doc, "60 条图发布", "results/auditor/interaction_graph_release_v1/")
    add_source_path(doc, "新版双人验证协议", "results/auditor/interaction_graph_double_validation_v1/PROTOCOL.md")
    add_source_path(doc, "新 30 条盲选说明", "results/auditor/interaction_graph_blind30_v2/README.md")
    add_source_path(doc, "Qwen 五折协议", "docs/interaction_graph_qwen_5fold_protocol.md")
    add_source_path(doc, "Qwen 五折结果", "results/auditor/interaction_graph_blind30_v2/qwen_5fold_v1/qwen/RESULTS.md")
    add_source_path(doc, "Qwen 评分明细", "results/auditor/interaction_graph_blind30_v2/qwen_5fold_v1/qwen/score/agreement.json")

    # Document-wide paragraph safety settings.
    for p in doc.paragraphs:
        p_pr = p._p.get_or_add_pPr()
        widow = p_pr.find(qn("w:widowControl"))
        if widow is None:
            widow = OxmlElement("w:widowControl")
            p_pr.append(widow)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
